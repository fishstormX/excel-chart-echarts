"""Generate sample workbooks covering the supported chart types (openpyxl + XlsxWriter)."""
import datetime as dt
import sys
from pathlib import Path

out = Path(sys.argv[1] if len(sys.argv) > 1 else "samples")
out.mkdir(exist_ok=True)

from openpyxl import Workbook
from openpyxl.chart import (AreaChart, BarChart, DoughnutChart, LineChart, PieChart, RadarChart, Reference,
                            ScatterChart, Series, BubbleChart)
from openpyxl.chart.label import DataLabelList

wb = Workbook()
ws = wb.active
ws.title = "Data"
rows = [["Quarter", "North", "South", "East"], ["Q1", 120, 90, 60], ["Q2", 150, 110, 80],
        ["Q3", 170, 95, 120], ["Q4", 210, 130, 140]]
for r in rows:
    ws.append(r)


def axes_on(ch):
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
    return ch


def place(ch, cell, title):
    ch.title = title
    ch.width, ch.height = 14, 7.5
    ws.add_chart(ch, cell)


data = Reference(ws, min_col=2, max_col=4, min_row=1, max_row=5)
cats = Reference(ws, min_col=1, min_row=2, max_row=5)
for i, (grouping, overlap) in enumerate([("clustered", None), ("stacked", 100), ("percentStacked", 100)]):
    c = axes_on(BarChart()); c.type = "col"; c.grouping = grouping
    if overlap:
        c.overlap = overlap
    c.add_data(data, titles_from_data=True); c.set_categories(cats)
    c.y_axis.title = "Sales"; c.x_axis.title = "Quarter"
    if grouping == "stacked":
        c.dataLabels = DataLabelList(); c.dataLabels.showVal = True
    place(c, f"F{1 + i * 16}", f"Column {grouping}")
c = axes_on(BarChart()); c.type = "bar"; c.add_data(data, titles_from_data=True); c.set_categories(cats)
c.legend.position = "b"; place(c, "P1", "Horizontal bar")
c = axes_on(LineChart()); c.add_data(data, titles_from_data=True); c.set_categories(cats)
c.series[1].smooth = True; c.series[2].marker.symbol = "triangle"
place(c, "P17", "Line (smooth / markers)")
c = axes_on(AreaChart()); c.grouping = "stacked"; c.add_data(data, titles_from_data=True); c.set_categories(cats)
place(c, "P33", "Stacked area")
c = PieChart(); c.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True); c.set_categories(cats)
c.dataLabels = DataLabelList(); c.dataLabels.showPercent = True
place(c, "Z1", "Pie (North share)")
c = DoughnutChart(); c.holeSize = 50; c.add_data(Reference(ws, min_col=3, min_row=1, max_row=5), titles_from_data=True)
c.set_categories(cats); place(c, "Z17", "Doughnut holeSize 50")
c = axes_on(RadarChart()); c.type = "marker"; c.add_data(data, titles_from_data=True); c.set_categories(cats)
place(c, "Z33", "Radar")

# scatter + bubble on its own sheet
ws2 = wb.create_sheet("XY")
ws2.append(["x", "y", "size"])
for x in range(1, 11):
    ws2.append([x, x * x / 3 + (x % 3), x * 10])
sc = axes_on(ScatterChart()); sc.style = 13
s = Series(Reference(ws2, min_col=2, min_row=2, max_row=11), Reference(ws2, min_col=1, min_row=2, max_row=11), title="y = f(x)")
s.marker.symbol = "circle"; s.graphicalProperties.line.noFill = True
sc.series.append(s); sc.title = "Scatter (markers only)"; sc.x_axis.title = "x"; sc.y_axis.title = "y"
ws2.add_chart(sc, "E2")
bc = axes_on(BubbleChart())
bs = Series(values=Reference(ws2, min_col=2, min_row=2, max_row=11), xvalues=Reference(ws2, min_col=1, min_row=2, max_row=11),
            zvalues=Reference(ws2, min_col=3, min_row=2, max_row=11), title="bubbles")
bc.series.append(bs); bc.title = "Bubble"; ws2.add_chart(bc, "E18")

# combo bar + line on secondary axis with formulas (no cached values → evaluator path)
ws3 = wb.create_sheet("Combo")
ws3.append(["Month", "Revenue", "Cost", "Margin %"])
for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun"], start=2):
    ws3.append([m, 100 + i * 20, 60 + i * 9, f"=(B{i}-C{i})/B{i}"])
    ws3[f"D{i}"].number_format = "0.0%"
bar = axes_on(BarChart()); bar.add_data(Reference(ws3, min_col=2, max_col=3, min_row=1, max_row=7), titles_from_data=True)
bar.set_categories(Reference(ws3, min_col=1, min_row=2, max_row=7)); bar.y_axis.title = "Amount"
line = LineChart(); line.add_data(Reference(ws3, min_col=4, min_row=1, max_row=7), titles_from_data=True)
line.y_axis.axId = 200; line.y_axis.title = "Margin"; line.y_axis.crosses = "max"; line.y_axis.delete = False
line.y_axis.majorGridlines = None; line.y_axis.number_format = "0%"
bar += line; bar.title = "Combo bar + line (secondary axis, formulas)"; bar.width, bar.height = 16, 8
ws3.add_chart(bar, "F2")

# date category axis
ws4 = wb.create_sheet("Dates")
ws4.append(["Date", "Visitors"])
d0 = dt.date(2026, 1, 1)
for i in range(12):
    ws4.append([d0 + dt.timedelta(days=i * 7), 300 + (i * 37) % 200])
    ws4.cell(i + 2, 1).number_format = "yyyy-mm-dd"
from openpyxl.chart.axis import DateAxis
lc = LineChart(); lc.title = "Weekly visitors (date axis)"
lc.add_data(Reference(ws4, min_col=2, min_row=1, max_row=13), titles_from_data=True)
lc.set_categories(Reference(ws4, min_col=1, min_row=2, max_row=13))
lc.x_axis = DateAxis(crossAx=100); lc.x_axis.number_format = "m/d"; lc.x_axis.majorTimeUnit = "days"; lc.x_axis.delete = False
lc.y_axis.crossAx = 500; lc.y_axis.delete = False
ws4.add_chart(lc, "D2")
wb.save(out / "sample_openpyxl.xlsx")

# ---------- XlsxWriter (writes cached values like Excel) ----------
import xlsxwriter
wb = xlsxwriter.Workbook(str(out / "sample_xlsxwriter.xlsx"))
ws = wb.add_worksheet("Sheet1")
date_fmt = wb.add_format({"num_format": "yyyy-mm-dd"})
ws.write_row(0, 0, ["Region", "2024", "2025", "Growth"])
regs = [("华东", 320, 410), ("华南", 280, 300), ("华北", 190, 260), ("西南", 120, 180), ("东北", 90, 85)]
for i, (r, a, b) in enumerate(regs, start=1):
    ws.write_row(i, 0, [r, a, b, (b - a) / a])
ch = wb.add_chart({"type": "column"})
for col, name in ((1, "2024"), (2, "2025")):
    ch.add_series({"name": ["Sheet1", 0, col], "categories": ["Sheet1", 1, 0, 5, 0], "values": ["Sheet1", 1, col, 5, col],
                   "data_labels": {"value": True}})
ln = wb.add_chart({"type": "line"})
ln.add_series({"name": "Growth", "categories": ["Sheet1", 1, 0, 5, 0], "values": ["Sheet1", 1, 3, 5, 3], "y2_axis": True,
               "marker": {"type": "circle", "size": 6}, "line": {"color": "#C00000", "width": 2.25}})
ln.set_y2_axis({"num_format": "0%", "name": "Growth"})
ch.combine(ln); ch.set_title({"name": "区域销售 (XlsxWriter combo)"}); ch.set_y_axis({"name": "万元", "major_gridlines": {"visible": True}})
ch.set_legend({"position": "bottom"}); ch.set_size({"width": 640, "height": 360})
ws.insert_chart("F2", ch)
pie = wb.add_chart({"type": "doughnut"})
pie.add_series({"name": "2025", "categories": ["Sheet1", 1, 0, 5, 0], "values": ["Sheet1", 1, 2, 5, 2],
                "data_labels": {"percentage": True, "category": True}, "points": [{"fill": {"color": "#FF6B6B"}}]})
pie.set_hole_size(60); pie.set_title({"name": "2025 占比"}); pie.set_rotation(90)
ws.insert_chart("F22", pie)
ws.write(8, 0, "Date"); ws.write(8, 1, "Close")
for i in range(10):
    ws.write_datetime(9 + i, 0, dt.datetime(2026, 3, 1) + dt.timedelta(days=i), date_fmt)
    ws.write(9 + i, 1, 50 + (i * 7) % 13)
ar = wb.add_chart({"type": "area"})
ar.add_series({"name": "Close", "categories": ["Sheet1", 9, 0, 18, 0], "values": ["Sheet1", 9, 1, 18, 1], "fill": {"color": "#70AD47"}})
ar.set_x_axis({"date_axis": True, "num_format": "mm-dd"}); ar.set_title({"name": "Area on date axis"}); ar.set_legend({"none": True})
ws.insert_chart("P2", ar)
hb = wb.add_chart({"type": "bar", "subtype": "stacked"})
for col in (1, 2):
    hb.add_series({"name": ["Sheet1", 0, col], "categories": ["Sheet1", 1, 0, 5, 0], "values": ["Sheet1", 1, col, 5, col]})
hb.set_title({"name": "Stacked horizontal bar"}); ws.insert_chart("P22", hb)
sc = wb.add_chart({"type": "scatter", "subtype": "smooth_with_markers"})
sc.add_series({"name": "curve", "categories": ["Sheet1", 1, 1, 5, 1], "values": ["Sheet1", 1, 2, 5, 2], "trendline": {"type": "linear"}})
sc.set_title({"name": "Scatter smooth + trendline"}); ws.insert_chart("F42", sc)
wb.close()
print("written to", out)
